import docx
from typing import Tuple, List

def extract_docx(docx_path: str) -> Tuple[str, List[str]]:
    """
    Extracts text from DOCX files, including paragraphs and tables.
    """
    review_flags: List[str] = []
    text_chunks: List[str] = []

    try:
        doc = docx.Document(docx_path)
        for p in doc.paragraphs:
            txt = p.text.strip()
            if txt:
                text_chunks.append(txt)

        # Also extract table text (often used in resumes for skills/experience matrices)
        for table in doc.tables:
            for row in table.rows:
                row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_cells:
                    text_chunks.append(" | ".join(row_cells))

    except Exception as e:
        return "", [f"docx_read_error: {str(e)}"]

    return "\n\n".join(text_chunks).strip(), review_flags
